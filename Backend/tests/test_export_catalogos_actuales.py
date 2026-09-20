"""Validación del export XLSX de catálogos (sin requerir DB si el archivo ya existe)."""

from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import load_workbook

BACKEND_ROOT = Path(__file__).resolve().parents[1]
XLSX_PATH = BACKEND_ROOT / "catalogos_digitaliza_revision.xlsx"
REQUIRED_SHEETS = [
    "RESUMEN",
    "Rubros",
    "Relevadores",
    "Inspectores",
    "Motivos_Notificacion",
    "Motivos_Comprobacion",
    "Juzgados",
    "Calles",
    "Distritos",
    "Contraproducencias",
    "Items_Inspeccion",
    "Tipos_Actuacion",
    "Turnos",
]


@pytest.mark.skipif(not XLSX_PATH.exists(), reason="Ejecutar scripts/export_catalogos_actuales.py primero")
def test_xlsx_exists_with_required_sheets():
    wb = load_workbook(XLSX_PATH, read_only=True)
    assert wb.sheetnames[:13] == REQUIRED_SHEETS


@pytest.mark.skipif(not XLSX_PATH.exists(), reason="Ejecutar scripts/export_catalogos_actuales.py primero")
def test_xlsx_row_counts_minimum():
    wb = load_workbook(XLSX_PATH, read_only=True)
    assert wb["Rubros"].max_row - 1 >= 1000
    assert wb["Calles"].max_row - 1 == 744
    assert wb["Juzgados"].max_row - 1 >= 800
    assert wb["Relevadores"].max_row - 1 >= 2
    wb.close()


@pytest.mark.skipif(not XLSX_PATH.exists(), reason="Ejecutar scripts/export_catalogos_actuales.py primero")
def test_rubros_has_review_columns():
    wb = load_workbook(XLSX_PATH, read_only=True)
    headers = [c.value for c in wb["Rubros"][1]]
    assert "accion_usuario" in headers
    assert "nombre_definitivo" in headers
    assert "clasificacion_actual" in headers
    wb.close()
