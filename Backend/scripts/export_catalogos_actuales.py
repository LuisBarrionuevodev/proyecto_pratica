#!/usr/bin/env python
"""
Exporta catálogos actuales de Digitaliza a XLSX para revisión manual (READ ONLY).

Uso:
  cd Backend
  python scripts/export_catalogos_actuales.py
  python scripts/export_catalogos_actuales.py --output catalogos_digitaliza_revision.xlsx

Solo ejecuta SELECT contra SQLALCHEMY_DATABASE_URI. No modifica la base.
"""

from __future__ import annotations

import argparse
import os
import re
import sys
from datetime import datetime
from pathlib import Path
from typing import Any, Iterable, Sequence
from urllib.parse import unquote

from dotenv import load_dotenv
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy import create_engine, text

BACKEND_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_OUTPUT = BACKEND_ROOT / "catalogos_digitaliza_revision.xlsx"

RUBROS_CANONICOS_SEED = frozenset(
    {
        "Comestibles",
        "Carnicería",
        "Drugstore",
        "Kiosco",
        "Supermercado",
        "Pollería",
        "Pescadería",
        "Bar",
        "Cervecería",
        "Rotisería",
        "Cafetería",
        "Verdulería",
        "Panadería",
    }
)

RUBRO_TEST_PATTERN = re.compile(
    r"(?:"
    r"Pr\d+Rub-|RubStab|Rub10B1_|^RubA_|^RubB_|^rubro-\d+|UniqRub|RubroEO_|"
    r"Notif\d+-|Comp\d+-|CarnPr\d+-|VerdPr\d+-|CafePr\d+-|RotiPr\d+-|"
    r"PollNum\d+-|PanPr\d+-|Carn\d+c[A-Z]?-|Notif\d+c-|Comp\d+c-|"
    r"^Rubro \d|Rub10B1_|^RubA_|^RubB_"
    r")",
    re.IGNORECASE,
)

RELEVADOR_TEST_PATTERN = re.compile(r"(QA|Inactivo-|test)", re.IGNORECASE)
JUZGADO_CANONICOS = frozenset({"JF1", "JF2", "JF3"})
JUZGADO_TEST_PATTERN = re.compile(
    r"(^JZRB|^JZSO|^JZED|^j_|doc test|Juzgado doc|Jz Rein|Jz Solo|Juz \d|JZCF|^JZ[0-9A-F]{4}$)",
    re.IGNORECASE,
)

HISTORICO_STRING_MSG = (
    "Comprobación guarda motivo como string; renombrar catálogo no modifica históricos automáticamente."
)
CONTRA_HISTORICO_MSG = (
    "Actuaciones guardan contraproducencia como string; renombrar catálogo no modifica históricos automáticamente."
)
TURNOS_OBS = "Frontend actualmente usa MANIANA/TARDE hardcoded."

REVIEW_HEADERS = frozenset(
    {
        "accion_usuario",
        "nombre_definitivo",
        "codigo_definitivo",
        "legajo_definitivo",
        "orden_definitivo",
        "observacion_usuario",
    }
)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Exporta catálogos a XLSX (read-only).")
    parser.add_argument(
        "--output",
        type=Path,
        default=DEFAULT_OUTPUT,
        help=f"Ruta del XLSX (default: {DEFAULT_OUTPUT})",
    )
    parser.add_argument(
        "--database-uri",
        default=None,
        help="URI MySQL (default: SQLALCHEMY_DATABASE_URI del .env)",
    )
    return parser.parse_args()


def database_uri_from_env(explicit: str | None) -> str:
    load_dotenv(BACKEND_ROOT / ".env")
    uri = (explicit or os.getenv("SQLALCHEMY_DATABASE_URI") or "").strip()
    if not uri:
        raise SystemExit("SQLALCHEMY_DATABASE_URI no configurada.")
    return uri


def database_name(uri: str) -> str:
    return unquote(uri.rsplit("/", 1)[-1].split("?", 1)[0])


def fetch_all(engine, sql: str, params: dict | None = None) -> list[dict[str, Any]]:
    with engine.connect() as conn:
        result = conn.execute(text(sql), params or {})
        return [dict(row._mapping) for row in result]


def classify_rubro(nombre: str) -> str:
    name = (nombre or "").strip()
    if name in RUBROS_CANONICOS_SEED:
        return "CANONICO_SEED"
    if RUBRO_TEST_PATTERN.search(name) or name.startswith("Rubro "):
        return "CONFIRMADO_TEST"
    if re.search(r"[0-9a-f]{6,}", name, re.IGNORECASE) and len(name) < 40:
        return "CONFIRMADO_TEST"
    if re.match(r"^[A-Za-zÁÉÍÓÚáéíóúñÑ][A-Za-zÁÉÍÓÚáéíóúñÑ0-9 /\-]{2,}$", name):
        return "PROBABLE_REAL"
    return "DUDOSO"


def classify_relevador(nombre: str) -> str:
    name = (nombre or "").strip()
    if name == "Fabian Esquivel":
        return "CANONICO_SEED"
    if RELEVADOR_TEST_PATTERN.search(name):
        return "CONFIRMADO_TEST"
    return "DUDOSO"


def classify_inspector(nombre: str, legajo: str) -> str:
    if RELEVADOR_TEST_PATTERN.search(nombre or ""):
        return "CONFIRMADO_TEST"
    if legajo and legajo.isdigit() and len(legajo) == 5:
        return "CANONICO_SEED"
    return "DUDOSO"


def classify_juzgado(codigo: str, nombre: str) -> str:
    code = (codigo or "").strip().upper()
    if code in JUZGADO_CANONICOS:
        return "CANONICO_ACTUAL"
    label = f"{codigo} {nombre}"
    if JUZGADO_TEST_PATTERN.search(label):
        return "CONFIRMADO_TEST"
    return "DUDOSO"


def classify_distrito(codigo: int | None, nombre: str) -> str:
    if codigo is not None and 1 <= int(codigo) <= 20 and (nombre or "").startswith("Distrito "):
        return "CANONICO_SEED"
    return "DUDOSO"


def empty_review_fields(extra: int = 0) -> list[Any]:
    return [""] * (3 + extra)


def sheet_rows_rubros(engine) -> list[list[Any]]:
    rows = fetch_all(
        engine,
        """
        SELECT r.id, r.nombre, r.created_at,
          (SELECT COUNT(*) FROM domicilio d WHERE d.rubro_id = r.id) AS domicilios_refs,
          (SELECT COUNT(*) FROM relevamiento rel WHERE rel.rubro_id = r.id) AS relevamientos_refs
        FROM rubro r
        ORDER BY relevamientos_refs DESC, domicilios_refs DESC, r.id
        """,
    )
    out: list[list[Any]] = []
    for r in rows:
        clasif = classify_rubro(r["nombre"])
        out.append(
            [
                r["id"],
                r["nombre"],
                int(r["domicilios_refs"] or 0),
                int(r["relevamientos_refs"] or 0),
                r["created_at"],
                clasif,
                "",
                "",
                "",
            ]
        )
    return out


def sheet_rows_relevadores(engine) -> list[list[Any]]:
    rows = fetch_all(
        engine,
        """
        SELECT r.id, r.nombre, r.activo,
          COUNT(rr.relevamiento_id) AS relevamientos_refs
        FROM relevador r
        LEFT JOIN relevamiento_relevador rr ON rr.relevador_id = r.id
        GROUP BY r.id, r.nombre, r.activo
        ORDER BY relevamientos_refs DESC, r.id
        """,
    )
    return [
        [
            r["id"],
            r["nombre"],
            bool(r["activo"]),
            int(r["relevamientos_refs"] or 0),
            classify_relevador(r["nombre"]),
            "",
            "",
            "",
        ]
        for r in rows
    ]


def sheet_rows_inspectores(engine) -> list[list[Any]]:
    rows = fetch_all(
        engine,
        """
        SELECT i.id, i.nombre, i.legajo, t.turno,
          (SELECT COUNT(*) FROM actuaciones_inspector ai WHERE ai.inspector_id = i.id) AS actuaciones_refs
        FROM inspector i
        LEFT JOIN turno t ON t.id = i.turno_id
        ORDER BY i.id
        """,
    )
    return [
        [
            r["id"],
            r["nombre"],
            r["legajo"],
            r["turno"],
            int(r["actuaciones_refs"] or 0),
            classify_inspector(r["nombre"], r["legajo"] or ""),
            "",
            "",
            "",
            "",
        ]
        for r in rows
    ]


def sheet_rows_motivos_notificacion(engine) -> list[list[Any]]:
    rows = fetch_all(
        engine,
        """
        SELECT m.id, m.nombre,
          (SELECT COUNT(*) FROM notificacion_motivo nm WHERE nm.motivo = m.id) AS notificaciones_refs
        FROM motivo m
        ORDER BY m.id
        """,
    )
    return [[r["id"], r["nombre"], int(r["notificaciones_refs"] or 0), "", "", ""] for r in rows]


def sheet_rows_motivos_comprobacion(engine) -> list[list[Any]]:
    rows = fetch_all(
        engine,
        """
        SELECT c.id, c.nombre,
          (SELECT COUNT(*) FROM comprobacion co WHERE co.motivo = c.nombre) AS uso_historico_aprox
        FROM catalog_motivo_comprobacion c
        ORDER BY c.id
        """,
    )
    return [
        [
            r["id"],
            r["nombre"],
            int(r["uso_historico_aprox"] or 0),
            "",
            "",
            "",
            HISTORICO_STRING_MSG,
        ]
        for r in rows
    ]


def sheet_rows_juzgados(engine) -> list[list[Any]]:
    rows = fetch_all(
        engine,
        """
        SELECT j.id, j.codigo, j.nombre,
          (SELECT COUNT(*) FROM oficio o WHERE o.juzgado_id = j.id) AS oficios_refs
        FROM juzgado_catalogo j
        ORDER BY CASE j.codigo WHEN 'JF1' THEN 1 WHEN 'JF2' THEN 2 WHEN 'JF3' THEN 3 ELSE 99 END, j.id
        """,
    )
    return [
        [
            r["id"],
            r["codigo"],
            r["nombre"],
            int(r["oficios_refs"] or 0),
            classify_juzgado(r["codigo"], r["nombre"]),
            "",
            "",
            "",
            "",
        ]
        for r in rows
    ]


def sheet_rows_calles(engine, dupe_bases: set[str]) -> list[list[Any]]:
    rows = fetch_all(
        engine,
        """
        SELECT id, nombre_canonico, canon_base, nombre_key, activo
        FROM calle_catalogo
        ORDER BY id
        """,
    )
    return [
        [
            r["id"],
            r["nombre_canonico"],
            r["canon_base"],
            r["nombre_key"],
            bool(r["activo"]),
            (r["canon_base"] or "").lower() in dupe_bases if r["canon_base"] else False,
            "",
            "",
            "",
        ]
        for r in rows
    ]


def sheet_rows_distritos(engine) -> list[list[Any]]:
    rows = fetch_all(
        engine,
        "SELECT id, codigo, nombre FROM distrito ORDER BY codigo",
    )
    return [
        [
            r["id"],
            r["codigo"],
            r["nombre"],
            classify_distrito(r["codigo"], r["nombre"]),
            "",
            "",
            "",
        ]
        for r in rows
    ]


def sheet_rows_contraproducencias(engine) -> list[list[Any]]:
    rows = fetch_all(
        engine,
        """
        SELECT c.id, c.nombre,
          (SELECT COUNT(*) FROM actuaciones a WHERE a.contraproducencia = c.nombre) AS uso_historico_aprox
        FROM catalog_contraproducencia c
        ORDER BY c.id
        """,
    )
    return [
        [
            r["id"],
            r["nombre"],
            int(r["uso_historico_aprox"] or 0),
            "",
            "",
            "",
            CONTRA_HISTORICO_MSG,
        ]
        for r in rows
    ]


def sheet_rows_items_inspeccion(engine) -> list[list[Any]]:
    rows = fetch_all(
        engine,
        "SELECT id, codigo, nombre, activo, orden, tipo_respuesta "
        "FROM item_acta_inspeccion ORDER BY orden, id",
    )
    return [
        [
            r["id"],
            r["codigo"],
            r["nombre"],
            bool(r["activo"]),
            r["orden"],
            r.get("tipo_respuesta") or "ESTADO",
            "",
            "",
            "",
            "",
        ]
        for r in rows
    ]


def sheet_rows_tipos_actuacion(engine) -> list[list[Any]]:
    rows = fetch_all(
        engine,
        """
        SELECT t.id, t.nombre,
          (SELECT COUNT(*) FROM actuaciones a WHERE a.tipo = t.nombre) AS uso_historico_aprox
        FROM catalog_tipo_actuacion t
        ORDER BY t.id
        """,
    )
    return [
        [r["id"], r["nombre"], int(r["uso_historico_aprox"] or 0), "", "", ""]
        for r in rows
    ]


def sheet_rows_turnos(engine) -> list[list[Any]]:
    rows = fetch_all(engine, "SELECT id, turno FROM turno ORDER BY id")
    return [[r["id"], r["turno"], TURNOS_OBS] for r in rows]


def count_by_classification(rows: Sequence[Sequence[Any]], class_index: int) -> dict[str, int]:
    counts: dict[str, int] = {}
    for row in rows:
        key = str(row[class_index])
        counts[key] = counts.get(key, 0) + 1
    return counts


def build_resumen(sheets_data: dict[str, dict[str, Any]]) -> list[list[Any]]:
    resumen: list[list[Any]] = []
    for catalogo, meta in sheets_data.items():
        rows = meta["rows"]
        clasif_idx = meta.get("clasif_index")
        clasif_counts = meta.get("clasif_counts") or {}
        confirmados = sum(
            clasif_counts.get(k, 0) for k in ("CONFIRMADO_TEST",) if clasif_counts
        )
        if not clasif_counts and clasif_idx is not None:
            confirmados = count_by_classification(rows, clasif_idx).get("CONFIRMADO_TEST", 0)
        canonicos = sum(
            clasif_counts.get(k, 0)
            for k in ("CANONICO_SEED", "CANONICO_ACTUAL")
            if clasif_counts
        )
        if not clasif_counts and clasif_idx is not None:
            c = count_by_classification(rows, clasif_idx)
            canonicos = c.get("CANONICO_SEED", 0) + c.get("CANONICO_ACTUAL", 0)
        necesita = meta.get("necesita_revision")
        if necesita is None:
            necesita = confirmados + (
                count_by_classification(rows, clasif_idx).get("DUDOSO", 0) if clasif_idx is not None else 0
            )
        resumen.append(
            [
                catalogo,
                len(rows),
                confirmados,
                canonicos,
                necesita,
                meta.get("fuente", "DB"),
                meta.get("observacion", ""),
            ]
        )
    return resumen


def autosize_columns(ws: Worksheet, headers: Sequence[str], rows: Sequence[Sequence[Any]]) -> None:
    for idx, header in enumerate(headers, start=1):
        max_len = len(str(header))
        for row in rows:
            if idx - 1 < len(row):
                max_len = max(max_len, len(str(row[idx - 1] or "")))
        ws.column_dimensions[get_column_letter(idx)].width = min(max(max_len + 2, 12), 60)


def write_sheet(
    wb: Workbook,
    title: str,
    headers: Sequence[str],
    rows: Sequence[Sequence[Any]],
) -> Worksheet:
    ws = wb.create_sheet(title=title)
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1F4E78")
    review_fill = PatternFill("solid", fgColor="FFF2CC")

    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        if header in REVIEW_HEADERS:
            cell.fill = PatternFill("solid", fgColor="C65911")
            cell.font = Font(bold=True, color="FFFFFF")

    for r_idx, row in enumerate(rows, start=2):
        for c_idx, value in enumerate(row, start=1):
            cell = ws.cell(row=r_idx, column=c_idx, value=value)
            header = headers[c_idx - 1]
            if header in REVIEW_HEADERS:
                cell.fill = review_fill

    ws.freeze_panes = "A2"
    if rows:
        last_col = get_column_letter(len(headers))
        last_row = len(rows) + 1
        table = Table(displayName=f"T_{title.replace(' ', '_')[:20]}", ref=f"A1:{last_col}{last_row}")
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        ws.add_table(table)
        ws.auto_filter.ref = f"A1:{last_col}{last_row}"

    autosize_columns(ws, headers, rows)
    return ws


def validate_workbook(path: Path, expected_sheets: dict[str, int]) -> None:
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True)
    names = wb.sheetnames
    for sheet, min_rows in expected_sheets.items():
        if sheet not in names:
            raise RuntimeError(f"Falta hoja: {sheet}")
        ws = wb[sheet]
        data_rows = max(ws.max_row - 1, 0)
        if data_rows < min_rows:
            raise RuntimeError(f"Hoja {sheet}: esperaba >= {min_rows} filas, tiene {data_rows}")
    wb.close()


def export_catalogos(output: Path, database_uri: str) -> dict[str, Any]:
    engine = create_engine(database_uri, pool_pre_ping=True)

    dupe_rows = fetch_all(
        engine,
        """
        SELECT canon_base FROM calle_catalogo
        WHERE canon_base IS NOT NULL AND canon_base != ''
        GROUP BY canon_base HAVING COUNT(*) > 1
        """,
    )
    dupe_bases = {(r["canon_base"] or "").lower() for r in dupe_rows}

    datasets: dict[str, tuple[list[str], list[list[Any]]]] = {
        "Rubros": (
            [
                "id",
                "nombre_actual",
                "domicilios_refs",
                "relevamientos_refs",
                "created_at",
                "clasificacion_actual",
                "accion_usuario",
                "nombre_definitivo",
                "observacion_usuario",
            ],
            sheet_rows_rubros(engine),
        ),
        "Relevadores": (
            [
                "id",
                "nombre_actual",
                "activo",
                "relevamientos_refs",
                "clasificacion_actual",
                "accion_usuario",
                "nombre_definitivo",
                "observacion_usuario",
            ],
            sheet_rows_relevadores(engine),
        ),
        "Inspectores": (
            [
                "id",
                "nombre",
                "legajo",
                "turno",
                "actuaciones_refs",
                "clasificacion_actual",
                "accion_usuario",
                "nombre_definitivo",
                "legajo_definitivo",
                "observacion_usuario",
            ],
            sheet_rows_inspectores(engine),
        ),
        "Motivos_Notificacion": (
            [
                "id",
                "nombre_actual",
                "notificaciones_refs",
                "accion_usuario",
                "nombre_definitivo",
                "observacion_usuario",
            ],
            sheet_rows_motivos_notificacion(engine),
        ),
        "Motivos_Comprobacion": (
            [
                "id",
                "nombre_actual",
                "uso_historico_aprox",
                "accion_usuario",
                "nombre_definitivo",
                "observacion_usuario",
                "advertencia_historica",
            ],
            sheet_rows_motivos_comprobacion(engine),
        ),
        "Juzgados": (
            [
                "id",
                "codigo",
                "nombre_actual",
                "oficios_refs",
                "clasificacion_actual",
                "accion_usuario",
                "codigo_definitivo",
                "nombre_definitivo",
                "observacion_usuario",
            ],
            sheet_rows_juzgados(engine),
        ),
        "Calles": (
            [
                "id",
                "nombre_canonico",
                "canon_base",
                "nombre_key",
                "activo",
                "posible_duplicado",
                "accion_usuario",
                "nombre_definitivo",
                "observacion_usuario",
            ],
            sheet_rows_calles(engine, dupe_bases),
        ),
        "Distritos": (
            [
                "id",
                "codigo",
                "nombre",
                "clasificacion_actual",
                "accion_usuario",
                "nombre_definitivo",
                "observacion_usuario",
            ],
            sheet_rows_distritos(engine),
        ),
        "Contraproducencias": (
            [
                "id",
                "nombre_actual",
                "uso_historico_aprox",
                "accion_usuario",
                "nombre_definitivo",
                "observacion_usuario",
                "advertencia_historica",
            ],
            sheet_rows_contraproducencias(engine),
        ),
        "Items_Inspeccion": (
            [
                "id",
                "codigo",
                "nombre",
                "activo",
                "orden",
                "tipo_respuesta",
                "accion_usuario",
                "nombre_definitivo",
                "orden_definitivo",
                "observacion_usuario",
            ],
            sheet_rows_items_inspeccion(engine),
        ),
        "Tipos_Actuacion": (
            [
                "id",
                "nombre_actual",
                "uso_historico_aprox",
                "accion_usuario",
                "nombre_definitivo",
                "observacion_usuario",
            ],
            sheet_rows_tipos_actuacion(engine),
        ),
        "Turnos": (
            ["id", "turno", "observacion"],
            sheet_rows_turnos(engine),
        ),
    }

    meta_for_resumen: dict[str, dict[str, Any]] = {}
    row_counts: dict[str, int] = {}

    rubros_rows = datasets["Rubros"][1]
    juzgados_rows = datasets["Juzgados"][1]
    rubros_clasif = count_by_classification(rubros_rows, 5)
    juzgados_clasif = count_by_classification(juzgados_rows, 4)

    meta_for_resumen["Rubros"] = {
        "rows": rubros_rows,
        "clasif_index": 5,
        "clasif_counts": rubros_clasif,
        "fuente": "DB + run.py seed",
        "observacion": "Alta contaminación test en nombres con sufijos/QA.",
    }
    meta_for_resumen["Relevadores"] = {
        "rows": datasets["Relevadores"][1],
        "clasif_index": 4,
        "fuente": "DB + relevadores_canonicos.py",
    }
    meta_for_resumen["Inspectores"] = {
        "rows": datasets["Inspectores"][1],
        "clasif_index": 5,
        "fuente": "DB + inspectores_canonicos.py",
    }
    meta_for_resumen["Motivos_Notificacion"] = {
        "rows": datasets["Motivos_Notificacion"][1],
        "necesita_revision": 0,
        "fuente": "DB + run.py seed",
    }
    meta_for_resumen["Motivos_Comprobacion"] = {
        "rows": datasets["Motivos_Comprobacion"][1],
        "necesita_revision": len(datasets["Motivos_Comprobacion"][1]),
        "fuente": "DB + run.py seed",
        "observacion": HISTORICO_STRING_MSG,
    }
    meta_for_resumen["Juzgados"] = {
        "rows": juzgados_rows,
        "clasif_index": 4,
        "clasif_counts": juzgados_clasif,
        "fuente": "DB + run.py seed (JF1-3)",
        "observacion": "Mayoría fixtures de tests.",
    }
    meta_for_resumen["Calles"] = {
        "rows": datasets["Calles"][1],
        "necesita_revision": sum(1 for r in datasets["Calles"][1] if r[5]),
        "fuente": "DB + CSV import",
    }
    meta_for_resumen["Distritos"] = {
        "rows": datasets["Distritos"][1],
        "clasif_index": 3,
        "fuente": "DB + distritos.geojson",
    }
    meta_for_resumen["Contraproducencias"] = {
        "rows": datasets["Contraproducencias"][1],
        "necesita_revision": len(datasets["Contraproducencias"][1]),
        "fuente": "DB + run.py seed",
        "observacion": CONTRA_HISTORICO_MSG,
    }
    meta_for_resumen["Items_Inspeccion"] = {
        "rows": datasets["Items_Inspeccion"][1],
        "necesita_revision": len(datasets["Items_Inspeccion"][1]),
        "fuente": "DB + migraciones Alembic",
    }
    meta_for_resumen["Tipos_Actuacion"] = {
        "rows": datasets["Tipos_Actuacion"][1],
        "necesita_revision": len(datasets["Tipos_Actuacion"][1]),
        "fuente": "DB + run.py seed",
    }
    meta_for_resumen["Turnos"] = {
        "rows": datasets["Turnos"][1],
        "necesita_revision": len(datasets["Turnos"][1]),
        "fuente": "DB; UI hardcoded",
        "observacion": TURNOS_OBS,
    }

    resumen_headers = [
        "Catalogo",
        "Total",
        "Confirmados_test",
        "Canonicos_actuales",
        "Necesita_revision_usuario",
        "Fuente_actual",
        "Observacion",
    ]
    resumen_rows = build_resumen(meta_for_resumen)

    wb = Workbook()
    wb.remove(wb.active)
    write_sheet(wb, "RESUMEN", resumen_headers, resumen_rows)
    for sheet_name, (headers, rows) in datasets.items():
        write_sheet(wb, sheet_name, headers, rows)
        row_counts[sheet_name] = len(rows)
    row_counts["RESUMEN"] = len(resumen_rows)

    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)

    validate_workbook(
        output,
        {
            "RESUMEN": 10,
            "Rubros": 1000,
            "Relevadores": 2,
            "Inspectores": 20,
            "Motivos_Notificacion": 3,
            "Motivos_Comprobacion": 10,
            "Juzgados": 800,
            "Calles": 744,
            "Distritos": 20,
            "Contraproducencias": 11,
            "Items_Inspeccion": 6,
            "Tipos_Actuacion": 6,
            "Turnos": 2,
        },
    )

    return {
        "output": str(output.resolve()),
        "database": database_name(database_uri),
        "row_counts": row_counts,
        "generated_at": datetime.now().isoformat(timespec="seconds"),
    }


def main() -> int:
    args = parse_args()
    uri = database_uri_from_env(args.database_uri)
    print(f"Export READ ONLY desde: {database_name(uri)}")
    result = export_catalogos(args.output.resolve(), uri)
    print(f"XLSX generado: {result['output']}")
    print("Filas por hoja:")
    for sheet, count in result["row_counts"].items():
        print(f"  {sheet}: {count}")
    print("Confirmación: solo SELECT ejecutado; cero writes en la base.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
