#!/usr/bin/env python
"""
PREDEPLOY-CLEANUP-REAL-WHITELIST.1-DIAG — cruce XLSX administrativo vs sandbox (READ ONLY).
"""
from __future__ import annotations

import json
import re
from collections import defaultdict
from datetime import datetime
from pathlib import Path
from typing import Any

from dotenv import load_dotenv
from openpyxl import load_workbook
from sqlalchemy import create_engine, text

BACKEND_ROOT = Path(__file__).resolve().parents[1]
XLSX_PATH = BACKEND_ROOT / "docs" / "Listado_OT_Notificacion_Inspeccion_Comprobacion_Oficio.xlsx"
OUTPUT_PATH = BACKEND_ROOT / "scripts" / "output" / "protected_operational_manifest_20260920.json"
REPORT_PATH = BACKEND_ROOT / "scripts" / "output" / "whitelist_cross_diag_20260920.json"

COL_OT = 0
COL_NOTIF = 1
COL_INSP = 2
COL_COMP = 3
COL_OFICIO = 4

PLACEHOLDER_VALUES = frozenset({"000", "0000"})
AMBIGUOUS_PATTERN = re.compile(r"[/\-]|^\*")

DB_FIELDS = {
    "orden_trabajo": ("orden_trabajo", "numero_acta"),
    "notificacion": ("notificacion", "numero_acta"),
    "inspeccion": ("inspeccion", "numero_acta"),
    "comprobacion": ("comprobacion", "numero_acta"),
    "oficio": ("oficio", "numero_oficio"),
}

TEST_ACTUACIONES_SQL = """
    SELECT a.id, a.tipo, a.fecha, a.orden_trabajo_id, a.notificacion_id, a.comprobacion_id,
           ot.numero_acta AS ot_numero, ot.anio AS ot_anio,
           d.calle
    FROM actuaciones a
    LEFT JOIN orden_trabajo ot ON ot.id = a.orden_trabajo_id
    LEFT JOIN domicilio d ON d.id = a.domicilio_id
    WHERE d.calle REGEXP '^[0-9a-f]{8}-[0-9a-f]{4}-'
       OR ot.numero_acta REGEXP '^[0-9A-F]{5}[A-F]$'
       OR a.id = 11450
"""


def is_pure_numeric(s: str) -> bool:
    return bool(re.fullmatch(r"\d+", s))


def normalize_numeric(raw: str) -> str | None:
    s = raw.strip()
    if not is_pure_numeric(s):
        return None
    return str(int(s))


def classify_cell(raw: Any, col_idx: int, row_ot: Any) -> dict[str, Any]:
    if raw is None or str(raw).strip() == "":
        return {"kind": "EMPTY", "raw_value": None, "normalized_value": None}

    raw_value = str(raw).strip()

    if raw_value in PLACEHOLDER_VALUES and col_idx in (COL_INSP, COL_OFICIO):
        return {
            "kind": "PLACEHOLDER",
            "raw_value": raw_value,
            "normalized_value": None,
            "reason": "placeholder_revision",
        }

    if AMBIGUOUS_PATTERN.search(raw_value) or (col_idx == COL_OT and len(raw_value) > 6):
        candidates = []
        for part in re.split(r"[/\-]", raw_value):
            part = part.strip().lstrip("*")
            if part and is_pure_numeric(part) and part not in PLACEHOLDER_VALUES:
                candidates.append({"raw": part, "normalized": normalize_numeric(part)})
        return {
            "kind": "AMBIGUOUS",
            "raw_value": raw_value,
            "normalized_value": None,
            "interpretation_candidates": candidates,
        }

    if is_pure_numeric(raw_value):
        return {
            "kind": "NUMERIC",
            "raw_value": raw_value,
            "normalized_value": normalize_numeric(raw_value),
        }

    return {
        "kind": "NON_NUMERIC",
        "raw_value": raw_value,
        "normalized_value": raw_value,
    }


def load_xlsx_rows() -> list[dict[str, Any]]:
    wb = load_workbook(XLSX_PATH, read_only=True, data_only=True)
    ws = wb["Listado"]
    rows: list[dict[str, Any]] = []
    for i, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        ot_cell = classify_cell(row[COL_OT], COL_OT, row[COL_OT])
        rows.append(
            {
                "source_row": i,
                "ot": ot_cell,
                "notificacion": classify_cell(row[COL_NOTIF], COL_NOTIF, row[COL_OT]),
                "inspeccion": classify_cell(row[COL_INSP], COL_INSP, row[COL_OT]),
                "comprobacion": classify_cell(row[COL_COMP], COL_COMP, row[COL_OT]),
                "oficio": classify_cell(row[COL_OFICIO], COL_OFICIO, row[COL_OT]),
            }
        )
    wb.close()
    return rows


def fetch_all(conn, sql: str, params: dict | None = None) -> list[dict[str, Any]]:
    return [dict(r._mapping) for r in conn.execute(text(sql), params or {})]


def build_lookup(conn, table: str, column: str) -> dict[str, list[dict[str, Any]]]:
    rows = fetch_all(conn, f"SELECT * FROM `{table}`")
    by_norm: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for r in rows:
        raw = str(r[column] or "").strip()
        norm = normalize_numeric(raw) if is_pure_numeric(raw) else raw
        if norm:
            by_norm.setdefault(norm, []).append(r)
        by_norm.setdefault(f"raw:{raw}", []).append(r)
    return by_norm


def match_records(
    cell: dict[str, Any],
    lookup: dict[str, list[dict[str, Any]]],
    column: str,
) -> dict[str, Any]:
    if cell["kind"] in ("EMPTY", "PLACEHOLDER", "AMBIGUOUS"):
        return {"status": "SKIPPED", "reason": cell["kind"], "matches": []}

    raw = cell["raw_value"]
    norm = cell.get("normalized_value")
    candidates: list[dict[str, Any]] = []

    if norm and norm in lookup:
        candidates.extend(lookup[norm])
    if raw and f"raw:{raw}" in lookup:
        for c in lookup[f"raw:{raw}"]:
            if c not in candidates:
                candidates.append(c)

    # dedupe by id
    seen = set()
    unique = []
    for c in candidates:
        cid = c["id"]
        if cid not in seen:
            seen.add(cid)
            unique.append(c)

    if not unique:
        return {"status": "NOT_FOUND", "raw_value": raw, "normalized_value": norm, "matches": []}

    exact = [c for c in unique if str(c[column]).strip() == str(raw)]
    normalized_only = [c for c in unique if c not in exact]

    if len(unique) > 1:
        status = "MULTIPLE_MATCHES"
    elif exact:
        status = "MATCH_EXACT"
    else:
        status = "MATCH_NORMALIZED"

    return {
        "status": status,
        "raw_value": raw,
        "normalized_value": norm,
        "matches": unique,
        "exact_count": len(exact),
        "normalized_only_count": len(normalized_only),
    }


def actuaciones_for_ot(conn, ot_id: int) -> list[dict]:
    return fetch_all(
        conn,
        "SELECT id, tipo, fecha, notificacion_id, comprobacion_id FROM actuaciones WHERE orden_trabajo_id = :oid",
        {"oid": ot_id},
    )


def actuaciones_for_notif(conn, nid: int) -> list[dict]:
    return fetch_all(
        conn,
        "SELECT id, tipo, fecha, orden_trabajo_id FROM actuaciones WHERE notificacion_id = :nid",
        {"nid": nid},
    )


def actuaciones_for_comp(conn, cid: int) -> list[dict]:
    return fetch_all(
        conn,
        "SELECT id, tipo, fecha, orden_trabajo_id FROM actuaciones WHERE comprobacion_id = :cid",
        {"cid": cid},
    )


def actuacion_for_insp(conn, insp_id: int) -> dict | None:
    rows = fetch_all(
        conn,
        """
        SELECT a.id, a.tipo, a.fecha, a.orden_trabajo_id, a.notificacion_id, a.comprobacion_id
        FROM inspeccion i
        JOIN actuaciones a ON a.id = i.actuacion_id
        WHERE i.id = :iid
        """,
        {"iid": insp_id},
    )
    return rows[0] if rows else None


def oficio_relations(conn, oficio_id: int) -> dict[str, Any]:
    ofi = fetch_all(conn, "SELECT * FROM oficio WHERE id = :id", {"id": oficio_id})
    if not ofi:
        return {}
    o = ofi[0]
    exp = fetch_all(conn, "SELECT id, numero_expediente, anio, comprobacion_id, oficio_id FROM expediente WHERE oficio_id = :oid", {"oid": oficio_id})
    comp_id = o.get("comprobacion_id")
    comp_acts = []
    if comp_id:
        comp_acts = actuaciones_for_comp(conn, comp_id)
    inis = fetch_all(conn, "SELECT id, tipo_iniciador, estado_iniciador FROM iniciador_ruta WHERE oficio_id = :oid", {"oid": oficio_id})
    return {
        "oficio": o,
        "expedientes": exp,
        "comprobacion_id": comp_id,
        "actuaciones_via_comprobacion": comp_acts,
        "iniciadores": inis,
    }


def ruta_context(conn, actuacion_id: int) -> dict[str, Any]:
    items = fetch_all(
        conn,
        """
        SELECT ri.id, ri.ruta_trabajo_id, ri.iniciador_ruta_id, ri.estado_ejecucion,
               rt.created_by_user_id AS ruta_created_by
        FROM ruta_item ri
        LEFT JOIN ruta_trabajo rt ON rt.id = ri.ruta_trabajo_id
        WHERE ri.actuacion_id = :aid
        """,
        {"aid": actuacion_id},
    )
    return {"ruta_items": items}


def main() -> None:
    load_dotenv(BACKEND_ROOT / ".env")
    import os

    uri = os.getenv("SQLALCHEMY_DATABASE_URI", "").strip()
    if not uri:
        raise SystemExit("SQLALCHEMY_DATABASE_URI no configurada")

    xlsx_rows = load_xlsx_rows()
    ambiguous: list[dict] = []

    for row in xlsx_rows:
        for field in ("ot", "notificacion", "inspeccion", "comprobacion", "oficio"):
            cell = row[field]
            if cell["kind"] in ("AMBIGUOUS", "PLACEHOLDER"):
                ambiguous.append(
                    {
                        "source_row": row["source_row"],
                        "field": field,
                        **cell,
                    }
                )

    engine = create_engine(uri)
    protected: dict[str, list[dict]] = {
        "orden_trabajo": [],
        "actuaciones": [],
        "inspeccion": [],
        "notificacion": [],
        "comprobacion": [],
        "oficio": [],
        "expediente": [],
    }
    protected_act_ids: set[int] = set()
    protected_ot_ids: set[int] = set()

    stats = {
        "xlsx_rows": len(xlsx_rows),
        "ot": defaultdict(int),
        "notificacion": defaultdict(int),
        "inspeccion": defaultdict(int),
        "comprobacion": defaultdict(int),
        "oficio": defaultdict(int),
    }

    cross_details = {
        "orden_trabajo": [],
        "notificacion": [],
        "inspeccion": [],
        "comprobacion": [],
        "oficio": [],
    }

    def add_protected(entity: str, entry: dict) -> None:
        protected[entity].append(entry)

    acts_by_ot_ids: set[int] = set()
    acts_by_notif_ids: set[int] = set()
    acts_by_insp_ids: set[int] = set()
    acts_by_comp_ids: set[int] = set()
    acts_by_oficio_ids: set[int] = set()

    def protect_actuacion(act: dict, source_row: int, evidence: str, extra: dict | None = None) -> None:
        aid = act["id"]
        if "OT" in evidence:
            acts_by_ot_ids.add(aid)
        if "NOTIFICACION" in evidence:
            acts_by_notif_ids.add(aid)
        if "INSPECCION" in evidence:
            acts_by_insp_ids.add(aid)
        if "COMPROBACION" in evidence:
            acts_by_comp_ids.add(aid)
        if "OFICIO" in evidence:
            acts_by_oficio_ids.add(aid)
        if aid in protected_act_ids:
            return
        protected_act_ids.add(aid)
        entry = {
            "id": aid,
            "source_row": source_row,
            "evidence": evidence,
            "relations": extra or {},
        }
        add_protected("actuaciones", entry)

    with engine.connect() as conn:
        lookups = {
            k: build_lookup(conn, table, col) for k, (table, col) in DB_FIELDS.items()
        }

        # --- OT cross ---
        for row in xlsx_rows:
            m = match_records(row["ot"], lookups["orden_trabajo"], "numero_acta")
            stats["ot"][m["status"]] += 1
            detail: dict[str, Any] = {
                "source_row": row["source_row"],
                "match": m,
                "actuaciones": [],
                "ruta_items": [],
            }
            if m["status"] not in ("SKIPPED", "NOT_FOUND"):
                for ot in m["matches"]:
                    acts = actuaciones_for_ot(conn, ot["id"])
                    detail["actuaciones"].extend(acts)
                    protected_ot_ids.add(ot["id"])
                    add_protected(
                        "orden_trabajo",
                        {
                            "id": ot["id"],
                            "numero_raw": m["raw_value"],
                            "numero_normalizado": m["normalized_value"],
                            "numero_db": ot["numero_acta"],
                            "anio": ot["anio"],
                            "source_row": row["source_row"],
                            "evidence": "REAL_ADMIN_XLSX",
                            "match_status": m["status"],
                        },
                    )
                    for act in acts:
                        protect_actuacion(
                            act,
                            row["source_row"],
                            "REAL_ADMIN_XLSX_OT",
                            {"orden_trabajo_id": ot["id"]},
                        )
                        detail["ruta_items"].append(ruta_context(conn, act["id"]))
            cross_details["orden_trabajo"].append(detail)

        # --- Notificacion ---
        for row in xlsx_rows:
            m = match_records(row["notificacion"], lookups["notificacion"], "numero_acta")
            stats["notificacion"][m["status"]] += 1
            detail = {"source_row": row["source_row"], "match": m, "actuaciones": []}
            if m["status"] not in ("SKIPPED", "NOT_FOUND"):
                for n in m["matches"]:
                    add_protected(
                        "notificacion",
                        {
                            "id": n["id"],
                            "numero_raw": m["raw_value"],
                            "numero_normalizado": m["normalized_value"],
                            "numero_db": n["numero_acta"],
                            "anio": n["anio"],
                            "source_row": row["source_row"],
                            "evidence": "REAL_ADMIN_XLSX",
                            "match_status": m["status"],
                        },
                    )
                    acts = actuaciones_for_notif(conn, n["id"])
                    detail["actuaciones"].extend(acts)
                    for act in acts:
                        protect_actuacion(
                            act,
                            row["source_row"],
                            "REAL_ADMIN_XLSX_NOTIFICACION",
                            {"notificacion_id": n["id"]},
                        )
            cross_details["notificacion"].append(detail)

        # --- Inspeccion ---
        for row in xlsx_rows:
            m = match_records(row["inspeccion"], lookups["inspeccion"], "numero_acta")
            stats["inspeccion"][m["status"]] += 1
            detail = {"source_row": row["source_row"], "match": m, "actuaciones": []}
            if m["status"] not in ("SKIPPED", "NOT_FOUND"):
                for insp in m["matches"]:
                    add_protected(
                        "inspeccion",
                        {
                            "id": insp["id"],
                            "numero_raw": m["raw_value"],
                            "numero_normalizado": m["normalized_value"],
                            "numero_db": insp["numero_acta"],
                            "anio": insp["anio"],
                            "source_row": row["source_row"],
                            "evidence": "REAL_ADMIN_XLSX",
                            "match_status": m["status"],
                        },
                    )
                    act = actuacion_for_insp(conn, insp["id"])
                    if act:
                        detail["actuaciones"].append(act)
                        protect_actuacion(
                            act,
                            row["source_row"],
                            "REAL_ADMIN_XLSX_INSPECCION",
                            {"inspeccion_id": insp["id"]},
                        )
            cross_details["inspeccion"].append(detail)

        # --- Comprobacion ---
        for row in xlsx_rows:
            m = match_records(row["comprobacion"], lookups["comprobacion"], "numero_acta")
            stats["comprobacion"][m["status"]] += 1
            detail = {"source_row": row["source_row"], "match": m, "actuaciones": []}
            if m["status"] not in ("SKIPPED", "NOT_FOUND"):
                for c in m["matches"]:
                    add_protected(
                        "comprobacion",
                        {
                            "id": c["id"],
                            "numero_raw": m["raw_value"],
                            "numero_normalizado": m["normalized_value"],
                            "numero_db": c["numero_acta"],
                            "anio": c["anio"],
                            "source_row": row["source_row"],
                            "evidence": "REAL_ADMIN_XLSX",
                            "match_status": m["status"],
                        },
                    )
                    acts = actuaciones_for_comp(conn, c["id"])
                    detail["actuaciones"].extend(acts)
                    for act in acts:
                        protect_actuacion(
                            act,
                            row["source_row"],
                            "REAL_ADMIN_XLSX_COMPROBACION",
                            {"comprobacion_id": c["id"]},
                        )
            cross_details["comprobacion"].append(detail)

        # --- Oficio ---
        for row in xlsx_rows:
            m = match_records(row["oficio"], lookups["oficio"], "numero_oficio")
            stats["oficio"][m["status"]] += 1
            detail = {"source_row": row["source_row"], "match": m, "relations": []}
            if m["status"] not in ("SKIPPED", "NOT_FOUND"):
                for o in m["matches"]:
                    rel = oficio_relations(conn, o["id"])
                    detail["relations"].append(rel)
                    add_protected(
                        "oficio",
                        {
                            "id": o["id"],
                            "numero_raw": m["raw_value"],
                            "numero_normalizado": m["normalized_value"],
                            "numero_db": o["numero_oficio"],
                            "anio": o["anio"],
                            "source_row": row["source_row"],
                            "evidence": "REAL_ADMIN_XLSX",
                            "match_status": m["status"],
                            "relations": {
                                "expediente_ids": [e["id"] for e in rel.get("expedientes", [])],
                                "comprobacion_id": rel.get("comprobacion_id"),
                                "iniciador_ids": [i["id"] for i in rel.get("iniciadores", [])],
                            },
                        },
                    )
                    for e in rel.get("expedientes", []):
                        add_protected(
                            "expediente",
                            {
                                "id": e["id"],
                                "numero_raw": m["raw_value"],
                                "numero_normalizado": m["normalized_value"],
                                "source_row": row["source_row"],
                                "evidence": "REAL_ADMIN_XLSX_OFICIO",
                            },
                        )
                    for act in rel.get("actuaciones_via_comprobacion", []):
                        protect_actuacion(
                            act,
                            row["source_row"],
                            "REAL_ADMIN_XLSX_OFICIO_COMPROBACION",
                            {"oficio_id": o["id"]},
                        )
            cross_details["oficio"].append(detail)

        # --- 542 test actuaciones cross ---
        test_acts = fetch_all(conn, TEST_ACTUACIONES_SQL)
        test_ids = {t["id"] for t in test_acts}

        conflicts = []
        test_no_real = []
        for t in test_acts:
            aid = t["id"]
            if aid in protected_act_ids:
                # gather whitelist evidence
                insp = fetch_all(
                    conn,
                    "SELECT numero_acta, anio FROM inspeccion WHERE actuacion_id = :aid",
                    {"aid": aid},
                )
                conflicts.append(
                    {
                        "actuacion_id": aid,
                        "ot": t.get("ot_numero"),
                        "ot_anio": t.get("ot_anio"),
                        "tipo": t.get("tipo"),
                        "fecha": str(t.get("fecha")),
                        "calle": t.get("calle"),
                        "evidencia_test_anterior": (
                            "calle_uuid_fixture"
                            if t.get("calle") and re.match(r"^[0-9a-f]{8}-", t["calle"] or "", re.I)
                            else "ot_unique_ot_numero_pattern"
                            if t.get("ot_numero") and re.match(r"^[0-9A-F]{5}[A-F]$", t["ot_numero"] or "", re.I)
                            else "qa_habilitacion_persistence_test"
                            if aid == 11450
                            else "mixed"
                        ),
                        "evidencia_real_xlsx": "PROTECTED_REAL via whitelist cross",
                        "inspeccion": insp,
                        "resolucion_recomendada": "CONFLICT_REVIEW — retirar de auto-delete; clasificar PROTECTED_REAL",
                    }
                )
            else:
                test_no_real.append(aid)

        # --- TEST_WRAPPER around real ---
        wrappers = []
        if protected_act_ids:
            ph = ",".join(str(i) for i in protected_act_ids)
            wrapper_rows = fetch_all(
                conn,
                f"""
                SELECT ri.id AS ruta_item_id, ri.ruta_trabajo_id, ri.actuacion_id,
                       ri.iniciador_ruta_id, rt.created_by_user_id,
                       u.email AS ruta_creator_email
                FROM ruta_item ri
                JOIN ruta_trabajo rt ON rt.id = ri.ruta_trabajo_id
                LEFT JOIN users u ON u.id = rt.created_by_user_id
                WHERE ri.actuacion_id IN ({ph})
                  AND (u.email LIKE '%@t.local' OR u.email LIKE '%@test.local')
                """,
            )
            wrappers = wrapper_rows

    # dedupe protected lists by id
    for key in protected:
        seen = set()
        deduped = []
        for item in protected[key]:
            iid = item["id"]
            if iid not in seen:
                seen.add(iid)
                deduped.append(item)
        protected[key] = deduped

    manifest = {
        "generated_at": datetime.now().isoformat(),
        "source_file": str(XLSX_PATH.name),
        "mode": "READ_ONLY_PROPOSAL",
        "writes_executed": False,
        "db_fields_used": DB_FIELDS,
        "entities": protected,
        "summary": {
            "protected_actuaciones_distinct": len(protected_act_ids),
            "protected_orden_trabajo": len(protected["orden_trabajo"]),
            "protected_notificacion": len(protected["notificacion"]),
            "protected_inspeccion": len(protected["inspeccion"]),
            "protected_comprobacion": len(protected["comprobacion"]),
            "protected_oficio": len(protected["oficio"]),
            "protected_expediente": len(protected["expediente"]),
        },
    }

    report = {
        "generated_at": datetime.now().isoformat(),
        "database": "digitaliza_sandbox",
        "writes_executed": False,
        "normalization_rules": {
            "numeric": "strip leading zeros via int(); 004353 -> 4353",
            "ambiguous": "cells with /, -, *, len>6 OT — not auto-normalized",
            "placeholders": list(PLACEHOLDER_VALUES),
        },
        "db_fields": DB_FIELDS,
        "statistics": {k: dict(v) if isinstance(v, defaultdict) else v for k, v in stats.items()},
        "AMBIGUOUS_SOURCE_VALUES": ambiguous,
        "confirmado_real": {
            "actuaciones_distinct": len(protected_act_ids),
            "por_ot": len(acts_by_ot_ids),
            "por_notificacion": len(acts_by_notif_ids),
            "por_inspeccion": len(acts_by_insp_ids),
            "por_comprobacion": len(acts_by_comp_ids),
            "por_oficio": len(acts_by_oficio_ids),
        },
        "test_actuaciones_542": {
            "total": len(test_ids),
            "TEST_CONFIRMED_NO_REAL_MATCH": len(test_no_real),
            "CONFLICT_REAL_WHITELIST": len(conflicts),
            "conflicts": conflicts,
            "safe_auto_delete_ids_sample": test_no_real[:50],
        },
        "TEST_WRAPPER_AROUND_REAL_DATA": wrappers,
        "cleanup_candidates_to_remove_from_auto_delete": {
            "actuacion_ids": [c["actuacion_id"] for c in conflicts],
            "count": len(conflicts),
        },
        "cross_sample": {
            "orden_trabajo_not_found": [
                d for d in cross_details["orden_trabajo"] if d["match"]["status"] == "NOT_FOUND"
            ][:20],
            "orden_trabajo_multiple": [
                d for d in cross_details["orden_trabajo"] if d["match"]["status"] == "MULTIPLE_MATCHES"
            ][:20],
            "oficio_matches": [d for d in cross_details["oficio"] if d["match"]["status"] not in ("SKIPPED", "NOT_FOUND")],
        },
        "classification_precedence": "PROTECTED_REAL > CONFIRMADO_TEST > PROBABLE_TEST > INDETERMINADO",
    }

    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    OUTPUT_PATH.write_text(json.dumps(manifest, indent=2, default=str), encoding="utf-8")
    REPORT_PATH.write_text(json.dumps(report, indent=2, default=str), encoding="utf-8")

    print(json.dumps({
        "manifest": str(OUTPUT_PATH),
        "report": str(REPORT_PATH),
        "protected_actuaciones": len(protected_act_ids),
        "conflicts_with_542": len(conflicts),
        "stats": report["statistics"],
        "ambiguous_count": len(ambiguous),
    }, indent=2))


if __name__ == "__main__":
    main()
