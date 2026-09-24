#!/usr/bin/env python
"""
Genera/enriquece manifests FASE 1 (SELECT + lectura archivos, sin DELETE).

- Enriquece protected manifest con metadata XLSX, unresolved, ambiguous.
- Genera cleanup_manifest_phase1_YYYYMMDD.json congelado.
"""
from __future__ import annotations

import argparse
import json
import sys
from datetime import datetime
from pathlib import Path

BACKEND_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(BACKEND_ROOT))

from dotenv import load_dotenv
from sqlalchemy import create_engine

from app.domains.predeploy_cleanup.manifest_io import file_sha256, load_manifest, manifest_sha256
from app.domains.predeploy_cleanup.planner import build_cleanup_manifest_from_db

DEFAULT_PROTECTED = BACKEND_ROOT / "scripts" / "output" / "protected_operational_manifest_20260920.json"
DEFAULT_WHITELIST_DIAG = BACKEND_ROOT / "scripts" / "output" / "whitelist_cross_diag_20260920.json"
DEFAULT_XLSX = BACKEND_ROOT / "docs" / "Listado_OT_Notificacion_Inspeccion_Comprobacion_Oficio.xlsx"
OUTPUT_DIR = BACKEND_ROOT / "scripts" / "output"


def parse_args() -> argparse.Namespace:
    p = argparse.ArgumentParser(description="Genera manifests FASE 1 cleanup.")
    p.add_argument("--protected-manifest", type=Path, default=DEFAULT_PROTECTED)
    p.add_argument("--whitelist-diag", type=Path, default=DEFAULT_WHITELIST_DIAG)
    p.add_argument("--xlsx", type=Path, default=DEFAULT_XLSX)
    p.add_argument("--database-uri", default=None)
    return p.parse_args()


def _unresolved_ot_from_xlsx(xlsx_path: Path, conn) -> list[dict]:
    """OTs del XLSX ausentes en sandbox."""
    from openpyxl import load_workbook
    from sqlalchemy import text

    def norm(raw: str) -> str | None:
        s = str(raw).strip()
        if not s.isdigit():
            return None
        return str(int(s))

    wb = load_workbook(xlsx_path, read_only=True, data_only=True)
    ws = wb["Listado"]
    unresolved = []
    for i, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        raw = row[0]
        if raw is None:
            continue
        raw_s = str(raw).strip()
        if "/" in raw_s or "-" in raw_s or raw_s.startswith("*") or len(raw_s) > 6:
            continue
        if not raw_s.isdigit():
            continue
        n = norm(raw_s)
        found = conn.execute(
            text(
                """
                SELECT id FROM orden_trabajo
                WHERE numero_acta = :raw
                   OR CAST(numero_acta AS UNSIGNED) = CAST(:norm AS UNSIGNED)
                LIMIT 1
                """
            ),
            {"raw": raw_s, "norm": n},
        ).fetchone()
        if not found:
            unresolved.append(
                {
                    "raw_value": raw_s,
                    "normalized_value": n,
                    "source_row": i,
                    "evidence": "REAL_ADMIN_XLSX_NOT_IMPORTED",
                }
            )
    wb.close()
    return unresolved


def enhance_protected(
    protected: dict,
    xlsx_path: Path,
    whitelist_diag: dict,
    conn=None,
) -> dict:
    """Agrega metadata sin alterar clasificaciones de entities."""
    enhanced = dict(protected)
    enhanced["manifest_version"] = "20260920.1"
    enhanced["protection_precedence"] = "PROTECTED_REAL > CONFIRMADO_TEST > PROBABLE_TEST > INDETERMINADO"
    enhanced["source_administrative"] = {
        "path": str(xlsx_path.relative_to(BACKEND_ROOT)),
        "filename": xlsx_path.name,
        "size_bytes": xlsx_path.stat().st_size,
        "sha256": file_sha256(xlsx_path),
        "sheet": "Listado",
        "operational_rows": whitelist_diag.get("statistics", {}).get("xlsx_rows", 1357),
    }
    ot_unresolved = _unresolved_ot_from_xlsx(xlsx_path, conn) if conn else []
    enhanced["protected_unresolved_identifiers"] = {
        "orden_trabajo_not_found": ot_unresolved,
        "orden_trabajo_not_found_count": len(ot_unresolved),
        "note": "IDs administrativos ausentes en sandbox; no participan en DELETE",
    }
    # Completar unresolved desde stats NOT_FOUND si hace falta
    enhanced["ambiguous_values"] = [
        {**item, "clasificacion": "REVIEW_REQUIRED"}
        for item in whitelist_diag.get("AMBIGUOUS_SOURCE_VALUES", [])
    ]
    enhanced["manifest_sha256"] = manifest_sha256(enhanced)
    return enhanced


def main() -> None:
    args = parse_args()
    load_dotenv(BACKEND_ROOT / ".env")
    import os

    uri = (args.database_uri or os.getenv("SQLALCHEMY_DATABASE_URI") or "").strip()
    if not uri:
        raise SystemExit("SQLALCHEMY_DATABASE_URI no configurada")

    protected = load_manifest(args.protected_manifest)
    whitelist_diag = json.loads(args.whitelist_diag.read_text(encoding="utf-8"))
    engine = create_engine(uri)
    with engine.connect() as conn:
        enhanced = enhance_protected(protected, args.xlsx, whitelist_diag, conn=conn)
        OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
        protected_out = OUTPUT_DIR / "protected_operational_manifest_20260920.json"
        protected_out.write_text(json.dumps(enhanced, indent=2, default=str), encoding="utf-8")
        cleanup = build_cleanup_manifest_from_db(conn, enhanced)
    cleanup["manifest_sha256"] = manifest_sha256(cleanup)
    cleanup_out = OUTPUT_DIR / "cleanup_manifest_phase1_20260920.json"
    cleanup_out.write_text(json.dumps(cleanup, indent=2, default=str), encoding="utf-8")

    print(
        json.dumps(
            {
                "protected_manifest": str(protected_out),
                "protected_sha256": enhanced["manifest_sha256"],
                "cleanup_manifest": str(cleanup_out),
                "cleanup_sha256": cleanup["manifest_sha256"],
                "cleanup_actuaciones": len(cleanup["entities"].get("actuaciones", [])),
                "writes_executed": False,
            },
            indent=2,
        )
    )


if __name__ == "__main__":
    main()
